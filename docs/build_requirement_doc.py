from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "政企客户欠费与到期管理工具需求规格说明书_R1.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, after=5, line=1.35, first_indent=True):
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if first_indent:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)


def add_text(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    style_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.left_indent = Cm(0.74 + level * 0.5)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.left_indent = Cm(0.74)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    for run in p.runs:
        set_run_font(run, size=14 if level == 1 else 11.5, bold=True)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, title in enumerate(headers):
        cell = header.cells[idx]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        set_run_font(r, size=9.5, bold=True, color="FFFFFF")
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cell = cells[idx]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F4F7FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if idx == 0 and len(row_values) <= 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=9.2)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("政企客户欠费与到期管理工具需求规格说明书")
    set_run_font(r, size=8.5, color="666666")


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    add_footer(section)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("政企客户欠费与到期管理工具需求规格说明书")
    set_run_font(run, size=20, bold=True)

    meta = [
        "版本：1.0（R1）",
        "日期：2026-10-09",
        "需求来源：政企客户经理及政企后台人员沟通记录",
        "文档定位：原始客户需求整理及后续方案设计依据",
    ]
    for line in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        set_run_font(r, size=10, color="555555")
    doc.add_paragraph()

    add_heading(doc, "1. 业务背景", 1)
    add_heading(doc, "1.1 业务现状", 2)
    add_text(doc, "政企客户经理目前需要同时关注客户余额、欠费、专线到期和续签等事项。相关信息主要来自短信、线下表格、其他台席下发的欠费明细表以及现有业务系统。")
    add_text(doc, "客户经理虽然会收到余额不足或欠费预警短信，但实际核查后发现，部分短信存在不准确的情况。同时，客户经理还会收到产品、集团及其他类型的大量信息，重要提醒容易被淹没，无法依靠短信完成持续、准确的欠费管控。")
    add_heading(doc, "1.2 当前管理方式", 2)
    for item in [
        "根据每月专线到期明细表开展续签工作，并在过程中关注客户余额。",
        "根据其他台席下发的欠费明细表，客户经理在开展欠费回收时同步处理部分欠费客户。",
        "由客户经理自行整理欠费、余额和到期台账，但大部分客户经理未形成稳定的台账维护习惯。",
    ]:
        add_number(doc, item)
    add_heading(doc, "1.3 现存问题", 2)
    for item in [
        "信息来源分散，数据需要人工汇总、比对和核实。",
        "短信预警准确性和可用性不足，无法直接作为唯一判断依据。",
        "信息量大，客户经理无法逐条关注，存在遗漏。",
        "欠费、余额、专线到期和续签进度缺少统一台账。",
        "任务主要依赖人员主动发现和跟进，缺少自动提醒、催办和闭环记录。",
        "部分客户在停机后才缴费，存在客户感知和业务连续性风险。",
        "后台同事需要反复维护线下表格、汇总数据和下发明细，人工工作量较大。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "1.4 建设目标", 2)
    add_text(doc, "建设一个面向政企客户欠费和到期管理的内部工具，整合可获取的系统数据和线下数据，形成统一台账，并围绕预警、任务、督办和处理结果建立闭环。")
    add_text(doc, "第一阶段重点解决以下问题：")
    for item in [
        "让客户经理能够快速看到真正需要处理的客户。",
        "减少人工整理表格、重复核对和重复下发明细的工作。",
        "通过规则化预警和任务提醒降低欠费遗漏风险。",
        "在客户停机前尽早发现并推动客户经理跟进。",
        "为后台人员提供整体欠费风险和任务处理情况的可视化依据。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2. 使用场景", 1)
    scenarios = [
        ("2.1 月度专线到期和续签管理", "每月导入或同步专线到期明细，系统按到期时间生成待办任务。客户经理查看所属客户的专线、余额、欠费和续签状态，跟进客户续签或缴费，并记录处理结果。"),
        ("2.2 欠费明细集中处理", "后台人员获取其他台席下发的欠费明细后，通过系统导入或同步。系统识别客户归属、欠费金额、欠费时长和当前处理状态，形成客户经理待办任务，减少通过群聊、邮件或多个表格反复下发的工作。"),
        ("2.3 余额不足和停机风险预警", "系统根据余额、欠费、业务状态、专线到期时间等信息，按照已确认的预警规则识别风险客户。对需要及时处理的客户生成预警，并向对应客户经理或后台人员发送内部提醒。"),
        ("2.4 客户经理跟进处理", "客户经理在任务列表中查看待处理事项，确认客户联系情况、缴费情况、续签情况和其他处理结果。系统记录跟进时间、处理人、处理状态、备注和必要附件。"),
        ("2.5 逾期任务督办", "对超过规定处理时限仍未完成的任务，系统标记为逾期，并按权限向客户经理、直属负责人或后台督办人员提示。督办规则、提醒频率和升级范围需后续确认。"),
        ("2.6 后台集中监控", "政企后台人员按区域、客户经理、业务类型、风险等级、任务状态和时间范围查看整体情况，识别高风险客户、未处理任务和数据异常，并输出明细或统计结果。"),
        ("2.7 数据异常核查", "当短信、线下明细和系统数据不一致时，后台人员能够查看数据来源、更新时间和核查结果，标记异常原因，并将异常记录反馈给数据提供方或业务负责人。"),
        ("2.8 后续扩展场景", "在欠费管理试点稳定后，可评估扩展到专线续签、合同到期、客户服务事项和其他需要台账及任务督办的政企管理工作。扩展内容不作为本版本默认交付范围。"),
    ]
    for heading, body in scenarios:
        add_heading(doc, heading, 2)
        add_text(doc, body)

    add_heading(doc, "3. 用户群体和权限", 1)
    add_heading(doc, "3.1 用户群体", 2)
    add_table(doc, ["用户角色", "主要职责"], [
        ["客户经理", "查看所属客户风险和待办任务，跟进客户缴费、续签及其他处理事项，更新处理结果"],
        ["政企后台人员", "导入或维护数据，配置和核查任务，查看全局情况，开展欠费回收和督办"],
        ["台席或数据提供人员", "提供欠费、余额、专线到期等明细，必要时查看数据接收和处理状态"],
        ["直属负责人或督办人员", "查看团队任务和逾期情况，进行提醒、督办和处理结果检查"],
        ["系统管理员", "维护用户、角色、数据权限、基础配置、日志和系统运行状态"],
    ], widths=[4.0, 11.5])
    add_heading(doc, "3.2 权限原则", 2)
    for item in [
        "按角色和数据归属控制访问范围。",
        "客户经理默认只能查看和处理本人负责或授权范围内的客户。",
        "政企后台人员可按授权范围查看和维护整体数据。",
        "直属负责人可查看所属团队的客户和任务，不默认拥有跨团队维护权限。",
        "系统管理员负责系统配置和权限管理，不因技术权限自动获得业务数据使用授权。",
        "所有新增、修改、删除、导入、导出、分派、催办和关闭操作均应进行权限校验并记录日志。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "3.3 初步权限矩阵", 2)
    add_table(doc, ["功能", "客户经理", "政企后台", "督办人员", "系统管理员"], [
        ["查看本人客户风险", "是", "是", "是", "按授权"],
        ["查看全量客户风险", "否", "按授权", "按团队", "按授权"],
        ["导入欠费或到期明细", "否", "是", "否", "是"],
        ["修改客户处理结果", "是", "是", "按授权", "是"],
        ["分派和转派任务", "否", "是", "是", "是"],
        ["催办和关闭任务", "处理本人任务", "是", "是", "是"],
        ["导出数据", "按授权", "是", "按授权", "是"],
        ["配置预警规则", "否", "按授权", "否", "是"],
        ["查看操作日志", "否", "按授权", "按授权", "是"],
    ], widths=[5.5, 2.4, 2.4, 2.4, 2.4])
    add_text(doc, "上述权限矩阵为初步建议，正式权限需结合公司现有账号体系和平台能力确认。")

    add_heading(doc, "4. 制度范围", 1)
    add_heading(doc, "4.1 纳入范围", 2)
    for item in [
        "政企客户余额、欠费和停机风险管理。",
        "专线到期及续签提醒管理。",
        "欠费明细、专线到期明细及相关数据的集中管理。",
        "客户经理待办任务、跟进记录和处理结果管理。",
        "逾期任务提醒、督办和闭环管理。",
        "后台人员的数据核查、异常标记、统计查询和明细导出。",
        "与公司内部AI工具、数据平台和业务系统的合规接入。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "4.2 暂不纳入范围", 2)
    for item in [
        "直接替代现有计费、账务、客户关系管理或工单系统。",
        "自动修改客户账务、自动扣款或自动执行停复机。",
        "仅依赖短信内容进行自动判断并直接触发业务动作。",
        "面向外部客户的公开查询或客户自助缴费门户。",
        "未经确认的跨部门数据打通、外部平台接入或公网部署。",
        "复杂的预测模型、自动催收话术生成和自动外呼。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "4.3 管理制度和业务规则要求", 2)
    for item in [
        "系统预警结果用于辅助管理，不替代现有账务和业务系统的正式数据。",
        "账务金额、业务状态和停机状态应以公司正式数据源为准。",
        "线下表格导入的数据必须保留来源、批次、导入人和导入时间。",
        "处理人应对任务结果负责，关键处理结果应能够追溯。",
        "预警规则、任务时限和升级规则应由业务负责人确认后配置。",
        "数据冲突和数据异常不得被系统静默覆盖，应保留核查记录。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5. 回答要求", 1)
    add_heading(doc, "5.1 回答内容范围", 2)
    for item in [
        "某客户当前是否存在欠费、余额不足或停机风险。",
        "某客户的专线到期时间、续签状态和待处理事项。",
        "某客户经理名下的待办、逾期任务和高风险客户。",
        "某时间范围内的欠费数量、金额、处理数量和未处理数量。",
        "某条数据的来源、更新时间和核查状态。",
        "某项任务的处理记录、当前责任人和下一步动作。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "5.2 回答准确性要求", 2)
    for item in [
        "回答应优先使用经过授权的内部数据源。",
        "回答应标明数据更新时间或数据批次，避免将过期数据表达为实时数据。",
        "当多个来源数据不一致时，应明确提示“数据存在差异”，并列出来源和核查状态。",
        "当系统无法确认时，不得编造金额、状态、时间或处理结论。",
        "对高风险事项应给出依据字段和建议核查动作。",
        "对涉及客户敏感信息的回答，应遵循用户数据权限，不得越权展示。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "5.3 预警和任务生成要求", 2)
    for item in [
        "预警应说明触发规则、触发时间、关联客户和责任人。",
        "任务应说明任务来源、处理时限、当前状态和完成条件。",
        "系统不得仅凭一条未经核实的短信自动认定客户欠费或直接执行停机类操作。",
        "预警应支持人工确认、误报标记、暂缓处理和关闭，并保留操作记录。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "5.4 异常回答要求", 2)
    add_text(doc, "当数据缺失、来源冲突、更新时间过久或无法匹配客户经理时，系统应明确提示异常类型，并建议后台人员补充或核查，不应将异常数据当作正常数据展示。")

    add_heading(doc, "6. 部署和安全要求", 1)
    add_heading(doc, "6.1 部署要求", 2)
    for item in [
        "原则上部署在公司内部网络、公司批准的AI平台或公司内部业务平台。",
        "不使用未经公司批准的外部AI服务处理客户、账务、欠费或联系方式等内部信息。",
        "优先复用公司现有账号认证、权限管理、消息通知、日志和数据接口能力。",
        "系统部署方式、运行环境、服务器位置和网络访问范围需经信息安全及平台管理人员确认。",
        "开发、测试和生产环境应隔离，测试数据应进行脱敏或使用模拟数据。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "6.2 数据安全要求", 2)
    for item in [
        "客户名称、客户编号、联系方式、账务金额、欠费信息和业务状态属于内部敏感数据，应按公司制度管理。",
        "数据传输和存储应采用公司认可的安全机制。",
        "不同角色只能访问授权范围内的数据。",
        "导出文件应进行权限控制，并记录导出人、导出时间、筛选条件和数据范围。",
        "对删除、覆盖、批量修改和批量导入等高风险操作应设置确认机制和日志。",
        "数据应保留来源、更新时间、批次和处理状态，支持追溯。",
        "日志中不得无必要记录完整的敏感信息。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "6.3 AI使用安全要求", 2)
    for item in [
        "AI仅可在公司批准的内部工具和平台中使用。",
        "AI回答必须受权限和数据范围控制。",
        "AI生成的预警解释、任务建议或摘要不得直接替代业务人员的最终判断。",
        "涉及金额、停机、合同和客户责任的关键结论，应能够追溯到具体数据来源。",
        "需明确数据是否允许进入模型上下文、是否留存以及留存周期。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "6.4 运行和审计要求", 2)
    for item in [
        "记录登录、查询、导入、导出、任务分派、任务催办、规则修改和权限变更等关键操作。",
        "对数据同步失败、导入失败、接口异常和任务生成失败提供可查询日志。",
        "重要任务和预警应具备失败重试或人工补处理机制。",
        "系统应提供基础运行状态和异常监控能力。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "7. 交付要求", 1)
    add_heading(doc, "7.1 第一阶段建议交付范围", 2)
    for item in [
        "欠费明细和专线到期明细的导入或数据接入。",
        "客户、业务、欠费、余额和到期信息的统一展示。",
        "客户经理、客户归属和数据权限管理。",
        "基于规则的风险预警。",
        "待办任务生成、分派、提醒、催办和关闭。",
        "客户经理跟进记录和处理结果登记。",
        "后台全局查询、统计和明细导出。",
        "数据来源、更新时间、批次和异常核查记录。",
        "用户权限、操作日志和基础运行监控。",
    ]:
        add_number(doc, item)
    add_heading(doc, "7.2 需求分析交付物", 2)
    for item in [
        "需求规格说明书。",
        "现状业务流程和目标业务流程。",
        "用户角色和权限矩阵。",
        "数据来源和字段清单。",
        "预警规则清单。",
        "任务状态和处理流程说明。",
        "页面和交互原型或页面清单。",
        "接口及数据接入方案。",
        "安全和部署方案。",
        "验收用例和试点方案。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "7.3 开发和测试交付物", 2)
    for item in [
        "前端页面及操作说明。",
        "后端服务、接口说明和数据字典。",
        "数据库设计或数据映射说明。",
        "配置项和权限配置说明。",
        "单元测试、接口测试和核心流程测试记录。",
        "测试环境部署说明。",
        "用户试用问题清单和修复记录。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "7.4 上线和运维交付物", 2)
    for item in [
        "部署手册。",
        "用户操作手册。",
        "管理员配置手册。",
        "数据导入模板和导入规则说明。",
        "预警规则配置说明。",
        "日志、监控和故障处理说明。",
        "版本发布记录。",
        "试点验收记录和遗留问题清单。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "7.5 验收要求", 2)
    for item in [
        "能够按指定数据源导入或获取欠费、余额和专线到期数据。",
        "能够根据客户归属生成对应客户经理的任务。",
        "能够按已确认规则生成预警，并展示触发依据。",
        "客户经理能够完成查看、跟进、备注、提交结果和关闭任务。",
        "后台能够查看全量或授权范围内任务，并识别逾期事项。",
        "导入、修改、分派、催办、关闭和导出等关键操作可追溯。",
        "权限限制、敏感数据保护和部署要求通过相关审核。",
        "试点用户能够使用核心流程完成实际业务处理。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "8. 待确认问题", 1)
    questions = {
        "8.1 数据来源和准确性": [
            "欠费、余额、停机状态和专线到期数据分别来自哪些系统或表格？",
            "各数据源的字段、更新频率、数据格式和数据责任人是什么？",
            "是否可以通过接口获取，还是第一阶段仅支持Excel或其他文件导入？",
            "哪个系统是账务金额、余额和停机状态的正式数据源？",
            "短信数据是否需要接入系统，还是仅作为人工核查参考？",
            "不同来源数据不一致时，采用什么优先级和核查流程？",
        ],
        "8.2 预警规则": [
            "余额低于多少时触发余额预警？",
            "欠费金额、欠费天数和业务状态分别设置什么预警条件？",
            "专线到期提前多少天生成提醒？",
            "是否需要按客户等级、业务类型或客户经理设置不同规则？",
            "是否分为一般、重要和紧急预警等级？",
            "预警触发后多久生成任务，任务处理时限是多少？",
            "是否允许客户经理暂缓处理，暂缓需要填写什么原因？",
        ],
        "8.3 任务和督办流程": [
            "任务由系统自动分派，还是由后台人员审核后分派？",
            "客户归属关系从哪个系统获取，发生变更时如何同步？",
            "任务状态是否统一为待处理、处理中、待核实、已完成、已关闭、逾期等？",
            "哪些处理结果可以直接关闭任务，哪些需要后台复核？",
            "逾期后由谁督办，是否需要逐级升级？",
            "提醒渠道使用公司内部平台、企业微信、邮件或其他方式？",
            "是否支持任务转派、批量处理、批量关闭和批量备注？",
        ],
        "8.4 用户和权限": [
            "使用哪些公司账号体系登录？",
            "客户经理、后台人员、台席、直属负责人和系统管理员的正式角色名称是什么？",
            "数据权限按客户经理、区域、部门、产品还是其他维度划分？",
            "是否允许跨团队查询和导出？",
            "哪些字段属于限制展示或脱敏字段？",
            "谁负责维护客户归属、预警规则和任务时限？",
        ],
        "8.5 平台和部署": [
            "公司现有可用的AI工具、内部平台、消息平台和数据接口有哪些？",
            "是否已有统一身份认证、权限中心、日志平台和文件存储能力？",
            "系统部署在现有平台还是新建内部应用？",
            "是否允许使用数据库、低代码平台、工作流平台或智能表格能力？",
            "开发、测试和生产环境由哪个团队提供和维护？",
            "数据留存、备份、审计和灾备要求是什么？",
        ],
        "8.6 试点和验收": [
            "首批试点选择哪些客户经理、区域或业务类型？",
            "试点数据覆盖多长时间、多少客户和多少条业务？",
            "试点阶段希望重点观察哪些指标，例如漏提醒数量、任务处理及时率、人工表格减少量等？",
            "由哪个部门或人员负责最终确认需求和验收？",
            "预期上线时间、迭代节奏和后续推广范围是什么？",
        ],
    }
    for heading, items in questions.items():
        add_heading(doc, heading, 2)
        for item in items:
            add_number(doc, item)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
