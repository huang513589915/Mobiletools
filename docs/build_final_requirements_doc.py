from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "政企客户欠费与到期管理工具最终需求设计说明书_R2.docx"

def set_font(run, size=10.2, bold=False, color="000000"):
    name = "Microsoft YaHei"
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    for key in ("eastAsia", "ascii", "hAnsi"):
        rpr.rFonts.set(qn(f"w:{key}"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)

def cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = tc_mar.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")

def table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "D9D9D9")

def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)

def pformat(p, after=4, line=1.25, indent=True):
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if indent:
        p.paragraph_format.first_line_indent = Cm(.74)

def para(doc, text, bold_prefix=None, after=4):
    p = doc.add_paragraph()
    pformat(p, after=after)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix); set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):]); set_font(r)
    else:
        r = p.add_run(text); set_font(r)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(.74)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.18
    r = p.add_run(text); set_font(r)
    return p

def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Cm(.74)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.18
    r = p.add_run(text); set_font(r)
    return p

def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    for r in p.runs:
        set_font(r, size=14 if level == 1 else 11.5, bold=True)
    return p

def add_table(doc, headers, rows, widths=None, size=8.8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    table_borders(t)
    repeat_header(t.rows[0])
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; cell_margins(c); shade(c, "1F4E78")
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h); set_font(r, size=size, bold=True, color="FFFFFF")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            c = cells[i]; c.text = ""; cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2: shade(c, "F4F7FA")
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.08
            if i == 0 and len(row) <= 5: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value)); set_font(r, size=size)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("政企客户欠费与到期管理工具最终需求设计说明书")
    set_font(r, size=8.2, color="666666")

def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.9); sec.bottom_margin = Cm(1.8); sec.left_margin = Cm(2.1); sec.right_margin = Cm(2.1)
    add_footer(sec)
    doc.styles["Normal"].font.name = "Microsoft YaHei"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.styles["Normal"].font.size = Pt(10.2)

    p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10)
    r = p.add_run("政企客户欠费与到期管理工具\n最终需求设计说明书"); set_font(r, size=20, bold=True)
    for x in ["版本：2.0（客户确认版 R2）", "日期：2026-10-09", "文档状态：待客户确认签字"]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
        r = p.add_run(x); set_font(r, size=10, color="555555")
    doc.add_paragraph()

    heading(doc, "1. 文档目的和确认范围")
    para(doc, "本文件用于确认“政企客户欠费与到期管理工具”第一阶段的业务目标、功能范围、用户权限、数据要求、交互要求、安全要求、交付内容和验收标准。")
    para(doc, "客户确认本文件后，本文件将作为后续原型细化、技术设计、开发、测试和验收的主要依据。未在本文件中明确确认的新增需求，应按变更流程评估，不自动纳入本阶段交付。")
    para(doc, "HTML Demo 用于展示页面结构和核心交互流程。Demo 中的客户名称、金额、日期、预警阈值、任务数量和处理状态均为示例数据，不作为正式业务数据或正式规则。")

    heading(doc, "2. 业务背景和建设目标")
    heading(doc, "2.1 现状", 2)
    para(doc, "客户经理需要同时关注客户余额、欠费、专线到期和续签事项。当前信息主要来自短信、线下表格、其他台席下发的欠费明细表和现有业务系统。")
    para(doc, "部分短信提醒经核查存在不准确情况，且客户经理同时接收大量产品、集团等其他信息，重要提醒容易被遗漏。当前管理主要依靠月度专线到期表、欠费明细表和客户经理自行整理台账，存在信息分散、重复核对、任务遗漏和停机后缴费等问题。")
    heading(doc, "2.2 建设目标", 2)
    para(doc, "建设公司内部使用的欠费与到期管理工具，整合可获取的系统数据和线下数据，形成统一客户台账，并建立“数据进入、风险识别、任务生成、客户跟进、督办提醒、结果闭环”的管理流程。")
    for x in ["集中展示客户余额、欠费、专线到期和处理状态。", "降低人工整理、比对和重复下发明细的工作量。", "通过规则化预警和任务提醒降低欠费遗漏风险。", "在客户停机前尽早发现风险并推动客户经理跟进。", "为后台人员提供风险、任务、逾期和异常数据的统一视图。"]:
        bullet(doc, x)

    heading(doc, "3. 第一阶段建设范围")
    heading(doc, "3.1 纳入范围", 2)
    for x in ["欠费明细和专线到期明细的导入或数据接入。", "客户、业务、欠费、余额、到期和客户归属的统一台账。", "客户经理、后台人员、督办人员和系统管理员的权限控制。", "余额不足、欠费、停机风险和专线到期预警。", "待办任务生成、分派、转派、提醒、催办、处理和关闭。", "客户经理跟进记录、缴费结果、续签结果和备注登记。", "后台风险监控、任务查询、统计和明细导出。", "数据异常标记、来源对比、核查结论和处理记录。", "操作日志、运行异常日志和基础运行监控。", "在公司批准的内部平台中提供受权限控制的查询或智能辅助能力。"]:
        numbered(doc, x)
    heading(doc, "3.2 暂不纳入范围", 2)
    for x in ["替代现有计费、账务、客户关系管理或工单系统。", "自动修改账务、自动扣款或自动执行停复机。", "仅凭短信内容自动认定欠费并直接触发业务动作。", "面向外部客户的公开查询或自助缴费门户。", "未经安全和平台评审确认的外部平台接入或公网部署。", "复杂预测模型、自动外呼和自动催收话术生成。"]:
        bullet(doc, x)

    heading(doc, "4. 用户角色和权限")
    add_table(doc, ["用户角色", "主要职责", "默认数据范围", "主要操作"], [
        ["客户经理", "跟进客户缴费、续签和风险事项", "本人负责或授权客户", "查看、跟进、备注、提交结果、申请暂缓"],
        ["政企后台人员", "数据维护、风险核查、任务管理和整体督办", "授权范围内的整体数据", "导入、核查、分派、转派、催办、复核、导出"],
        ["台席或数据提供人员", "提供并跟踪欠费、余额、到期数据", "授权数据批次", "提供数据、查看接收状态、处理导入异常"],
        ["直属负责人或督办人员", "查看团队任务和逾期事项", "所属团队数据", "查看、催办、升级、检查处理结果"],
        ["系统管理员", "账号、权限、规则、日志和运行配置", "按公司授权", "用户管理、权限配置、规则配置、日志查看"],
    ], widths=[3.0, 5.0, 4.0, 5.0], size=8.4)
    para(doc, "权限原则：数据权限和功能权限分别控制；客户经理默认不能查看跨团队客户；导入、导出、批量修改、规则修改和关闭任务等操作必须进行权限校验；所有关键操作必须记录操作人、时间、对象和结果。")

    heading(doc, "5. 业务流程设计")
    heading(doc, "5.1 主流程", 2)
    for x in ["后台人员通过接口或模板导入欠费、余额、专线到期和客户归属数据。", "系统执行字段、金额、日期、重复数据和关联关系预检。", "系统保存数据来源、批次、更新时间和异常信息。", "系统依据已确认的预警规则识别风险客户。", "系统生成任务并按客户归属分派给客户经理，或进入后台待分派队列。", "客户经理查看任务，联系客户并登记缴费、续签、核实或其他处理结果。", "对临近时限或已逾期任务发送提醒、催办或升级通知。", "任务完成后关闭或提交后台复核，系统保留全过程记录。", "后台人员查看风险、任务、逾期和异常数据，并导出授权范围内的结果。"]:
        numbered(doc, x)
    heading(doc, "5.2 任务状态", 2)
    para(doc, "建议状态为：待处理、处理中、待核实、已完成、已关闭、已暂缓、已逾期。正式状态名称、状态转换条件和关闭条件以客户确认的业务规则为准。")
    heading(doc, "5.3 数据异常处理", 2)
    para(doc, "当系统数据、短信信息和线下明细不一致时，不直接静默覆盖。系统应保留各来源、批次和更新时间，后台人员标记异常原因、核查结论和后续责任人。")

    heading(doc, "6. 功能需求")
    modules = [
        ("6.1 风险工作台", "让客户经理和后台人员优先看到需要处理的客户和任务。", ["展示高风险客户、待处理任务、欠费金额和专线到期等指标。", "支持按全部、已逾期、专线到期、待核实等维度切换。", "支持按客户名称、客户编号、责任人、风险等级和处理状态筛选。", "展示客户、业务、风险、欠费金额、余额、到期时间、责任人和处理状态。", "支持进入客户详情、创建跟进任务和标记待核实。"]),
        ("6.2 客户统一台账", "统一查看客户、业务、欠费、余额、到期、续签和归属信息。", ["支持多条件查询。", "展示数据来源、更新时间、批次和异常标识。", "支持查看客户历史预警、任务和跟进记录。", "按权限控制查询和导出范围。"]),
        ("6.3 数据接入和导入预检", "将欠费、到期和客户归属数据可靠地导入或接入。", ["支持标准模板和接口接入。", "校验必填项、金额、日期、重复数据、客户匹配和业务关联。", "展示成功、失败、重复和待核实数量。", "支持错误明细下载并保存批次信息。"]),
        ("6.4 风险预警", "依据确认后的规则生成可解释的风险预警。", ["支持余额、欠费、停机风险和专线到期预警。", "展示触发规则、字段、时间、风险等级和责任人。", "支持确认、误报、暂缓、解除和重新触发。", "不得仅凭短信直接执行停机或账务动作。"]),
        ("6.5 任务中心", "将预警和到期事项转化为可执行、可追踪的任务。", ["支持自动分派、后台分派、转派和重新指派。", "支持任务状态流转。", "支持联系情况、缴费结果、续签结果、备注和附件。", "支持关闭前复核。"]),
        ("6.6 提醒、催办和督办", "减少任务遗漏并推动逾期事项处理。", ["支持创建、临近截止和逾期提醒。", "支持团队逾期查看、催办和升级。", "消息失败可查询并人工补处理。"]),
        ("6.7 后台监控和统计", "为后台提供整体风险和处理情况视图。", ["展示风险客户、欠费金额、待办、逾期、已完成和异常指标。", "支持筛选、明细、汇总和导出。", "统计结果可与明细数据核对。"]),
        ("6.8 数据异常核查", "处理多来源数据差异并保证可追溯。", ["识别缺失、重复、冲突、过期和无法匹配数据。", "支持来源对比、异常原因、核查结论和处理历史。"]),
        ("6.9 查询和智能辅助", "在批准的内部平台中提供受权限控制的查询和辅助回答。", ["回答标明更新时间或批次。", "数据冲突时提示差异和来源。", "无法确认时不得编造。", "不得越权展示敏感数据，AI不替代业务最终判断。"]),
    ]
    for title, goal, items in modules:
        heading(doc, title, 2)
        para(doc, "目标：" + goal, "目标：")
        para(doc, "功能要求：", "功能要求：")
        for x in items: bullet(doc, x)

    heading(doc, "7. Demo 对应关系")
    add_table(doc, ["Demo页面", "对应需求"], [
        ["风险工作台", "风险指标、客户筛选、重点客户、风险详情"],
        ["任务中心", "任务列表、处理任务、完成任务"],
        ["客户台账", "客户与业务统一台账"],
        ["数据接入", "导入批次和数据状态展示"],
        ["规则与权限", "预警规则示例、角色权限示例"],
        ["客户详情弹窗", "风险依据、数据来源、建议动作和任务创建"],
    ], widths=[4.5, 13.0], size=9)
    para(doc, "Demo 仅用于确认页面结构、信息层级和交互方向，不代表最终视觉稿、正式数据和正式接口。")

    heading(doc, "8. 数据、部署和安全要求")
    heading(doc, "8.1 数据要求", 2)
    for x in ["欠费、余额、停机状态和专线到期数据应明确正式来源。", "线下导入数据必须保留来源、批次、导入人和导入时间。", "客户和业务应具备可用于关联的唯一标识。", "数据缺失、冲突、重复和过期情况不得被静默忽略。"]:
        bullet(doc, x)
    heading(doc, "8.2 部署要求", 2)
    for x in ["原则上部署在公司内部网络、公司批准的AI平台或公司内部业务平台。", "开发、测试和生产环境隔离。", "测试环境使用脱敏数据或模拟数据。", "优先复用公司统一认证、权限、消息、日志、文件存储和数据接口能力。"]:
        bullet(doc, x)
    heading(doc, "8.3 安全和审计要求", 2)
    for x in ["敏感数据按公司制度管理并采用认可的传输和存储安全机制。", "导出文件受权限控制，并记录导出人、时间、条件和范围。", "记录登录、查询、导入、导出、规则修改、任务分派、催办、关闭和权限变更。", "日志不得无必要记录完整敏感信息。", "生产数据是否进入AI上下文、是否留存及留存周期，须经安全评审确认。"]:
        bullet(doc, x)

    heading(doc, "9. 交付物和验收标准")
    heading(doc, "9.1 交付物", 2)
    for x in ["最终需求设计说明书。", "业务流程和任务状态说明。", "用户角色和权限矩阵。", "数据源、字段和数据映射清单。", "预警规则和任务规则清单。", "页面原型或页面实现。", "接口说明、数据字典和配置说明。", "测试计划、测试用例和测试记录。", "部署手册、用户操作手册和管理员手册。", "试点验收报告和遗留问题清单。"]:
        bullet(doc, x)
    heading(doc, "9.2 第一阶段验收标准", 2)
    for x in ["能够按确认的数据源导入或获取欠费、余额和专线到期数据。", "数据保留来源、批次、更新时间和导入结果。", "能够根据客户归属生成客户经理任务。", "能够按确认规则生成预警，并展示触发依据。", "客户经理能够查看、跟进、备注、提交结果和关闭任务。", "后台能够查看授权范围的风险、任务、逾期和异常数据。", "支持提醒、催办和逾期识别。", "关键操作可追溯。", "权限、敏感数据保护、部署和安全要求通过审核。", "试点用户能够完成至少一轮真实业务处理。"]:
        numbered(doc, x)

    heading(doc, "10. 本次确认后仍需补充的业务参数")
    add_table(doc, ["参数类别", "待确认内容", "确认责任方"], [
        ["数据来源", "欠费、余额、停机、到期和客户归属正式来源", "业务负责人、数据负责人"],
        ["数据接入", "接口、Excel、人工上传或组合方式", "平台负责人、数据负责人"],
        ["预警阈值", "余额、欠费金额、欠费天数、到期提前天数", "业务负责人"],
        ["任务规则", "分派、时限、逾期升级、复核和关闭条件", "业务负责人"],
        ["提醒渠道", "内部平台、企业微信、邮件或其他渠道", "平台负责人"],
        ["权限范围", "角色名称、数据范围、导出和脱敏要求", "业务负责人、平台负责人"],
        ["安全策略", "数据留存、备份、审计和AI上下文策略", "安全负责人"],
        ["试点范围", "客户经理、区域、客户量、数据周期和试点时间", "业务负责人"],
    ], widths=[3.2, 10.0, 4.3], size=8.8)

    heading(doc, "11. 需求变更原则")
    para(doc, "客户确认后，新增功能、扩大数据范围、增加外部系统接入、修改核心流程或调整验收标准，应形成需求变更记录，说明变更原因、影响范围、工期影响和双方确认结果。")

    heading(doc, "12. 客户确认签字")
    para(doc, "客户确认以下内容：")
    for x in ["已阅读并理解本需求设计说明书。", "认可第一阶段建设目标、功能范围、暂不纳入范围和验收标准。", "认可 HTML Demo 仅作为交互原型，不代表正式数据和正式业务规则。", "同意将本文件作为后续详细设计、开发、测试和验收的依据。", "对第10节所列业务参数，在实施前按约定完成补充确认。"]:
        bullet(doc, x)
    add_table(doc, ["确认方", "单位/部门", "姓名", "职务", "签字", "日期"], [
        ["客户业务负责人", "", "", "", "", ""],
        ["客户技术/平台负责人", "", "", "", "", ""],
        ["客户信息安全负责人", "", "", "", "", ""],
        ["项目实施方负责人", "", "", "", "", ""],
    ], widths=[4.0, 4.2, 2.5, 3.0, 3.0, 2.5], size=8.8)
    para(doc, "客户意见或补充说明：")
    for _ in range(3):
        para(doc, "______________________________________________________________________________", after=8)
    doc.save(OUT)

if __name__ == "__main__":
    build()
